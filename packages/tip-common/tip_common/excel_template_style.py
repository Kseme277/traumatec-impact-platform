"""Mise en forme Excel TIP pour modèles d'import téléchargeables."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

# Palette TIP (alignée frontend brand)
BRAND = "3641F5"
BRAND_DARK = "252DAE"
HEADER_FG = "FFFFFF"
EXAMPLE_FILL = "ECF3FF"
ALT_ROW = "F8FAFC"
MUTED = "64748B"
GRID = "E2E8F0"
SUCCESS = "12B76A"

_thin = Border(
    left=Side(style="thin", color=GRID),
    right=Side(style="thin", color=GRID),
    top=Side(style="thin", color=GRID),
    bottom=Side(style="thin", color=GRID),
)


def style_import_workbook(
    ws: Worksheet,
    *,
    title: str,
    subtitle: str,
    headers: list[str],
    example_rows: list[list[object]],
    column_widths: dict[str, float] | None = None,
    date_columns: set[str] | None = None,
    number_columns: set[str] | None = None,
    percent_columns: set[str] | None = None,
    validations: dict[str, list[str]] | None = None,
    help_rows: list[tuple[str, str]] | None = None,
) -> None:
    """Construit une feuille d'import soignée (bandeau + en-têtes + exemples + filtres)."""
    date_columns = date_columns or set()
    number_columns = number_columns or set()
    percent_columns = percent_columns or set()
    validations = validations or {}
    help_rows = help_rows or []

    # Bandeau titre (ligne 1)
    last_col = get_column_letter(len(headers))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    banner = ws.cell(1, 1, title)
    banner.font = Font(name="Calibri", bold=True, size=16, color=HEADER_FG)
    banner.fill = PatternFill("solid", fgColor=BRAND)
    banner.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32

    # Sous-titre (ligne 2)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sub = ws.cell(2, 1, subtitle)
    sub.font = Font(name="Calibri", size=10, color=MUTED, italic=True)
    sub.fill = PatternFill("solid", fgColor="F2F7FF")
    sub.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 22

    # En-têtes (ligne 3) — style proche Projects.xlsx (fond gris/bleu, texte blanc)
    header_fill = PatternFill("solid", fgColor=BRAND_DARK)
    header_font = Font(name="Calibri", bold=True, size=11, color=HEADER_FG)
    for idx, name in enumerate(headers, start=1):
        cell = ws.cell(3, idx, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin
    ws.row_dimensions[3].height = 28

    # Lignes d'exemple
    example_fill = PatternFill("solid", fgColor=EXAMPLE_FILL)
    for row_offset, values in enumerate(example_rows):
        excel_row = 4 + row_offset
        for col_idx, value in enumerate(values, start=1):
            header = headers[col_idx - 1]
            cell = ws.cell(excel_row, col_idx, value)
            cell.font = Font(name="Calibri", size=11, color="1E293B")
            cell.fill = example_fill
            cell.border = _thin
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if header in date_columns:
                cell.number_format = "YYYY-MM-DD"
            elif header in percent_columns:
                cell.number_format = "0%"
            elif header in number_columns:
                cell.number_format = "#,##0.00"

    data_start = 4
    data_end = max(4 + len(example_rows), 50)

    # Validations (listes déroulantes)
    for header, choices in validations.items():
        if header not in headers:
            continue
        col = headers.index(header) + 1
        letter = get_column_letter(col)
        formula = '"' + ",".join(choices) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = "Choisissez une valeur de la liste"
        dv.errorTitle = "Valeur invalide"
        dv.prompt = "Sélectionnez une option"
        dv.promptTitle = header
        ws.add_data_validation(dv)
        dv.add(f"{letter}{data_start}:{letter}{data_end}")

    # Largeurs de colonnes
    for idx, name in enumerate(headers, start=1):
        width = (column_widths or {}).get(name)
        if width is None:
            width = max(12, min(36, len(name) + 4))
        ws.column_dimensions[get_column_letter(idx)].width = width

    # Filtres + freeze (bandeau + sous-titre + header)
    ws.auto_filter.ref = f"A3:{last_col}{data_end}"
    ws.freeze_panes = "A4"

    # Feuille Aide
    wb = ws.parent
    if "Aide" in wb.sheetnames:
        del wb["Aide"]
    help_ws = wb.create_sheet("Aide", 1)
    help_ws["A1"] = "Guide d'utilisation — Traumatec Impact Platform"
    help_ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=HEADER_FG)
    help_ws["A1"].fill = PatternFill("solid", fgColor=BRAND)
    help_ws.merge_cells("A1:B1")
    help_ws.row_dimensions[1].height = 28

    help_ws["A2"] = "Colonne"
    help_ws["B2"] = "Description"
    for cell in (help_ws["A2"], help_ws["B2"]):
        cell.font = Font(name="Calibri", bold=True, color=HEADER_FG)
        cell.fill = PatternFill("solid", fgColor=BRAND_DARK)
        cell.border = _thin

    for i, (col, desc) in enumerate(help_rows or [(h, "") for h in headers], start=3):
        help_ws.cell(i, 1, col).font = Font(name="Calibri", bold=True)
        help_ws.cell(i, 2, desc).font = Font(name="Calibri")
        for c in range(1, 3):
            help_ws.cell(i, c).border = _thin
            if i % 2 == 0:
                help_ws.cell(i, c).fill = PatternFill("solid", fgColor=ALT_ROW)

    help_ws.column_dimensions["A"].width = 32
    help_ws.column_dimensions["B"].width = 72
    help_ws.freeze_panes = "A3"
