"""Analyse IA (NVIDIA) des champs à remplacer dans chaque fichier paquet AO."""

from __future__ import annotations

import json
import logging
import os
import re
import zipfile
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from xml.etree import ElementTree as ET

from tip_common.french_label_patterns import APOSTROPHE, EVENEMENT

logger = logging.getLogger(__name__)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

CONTEXT_KEYS = (
    "project_number",
    "title",
    "city",
    "country",
    "lieu",
    "region",
    "responsible_person",
    "responsible_email",
    "responsible_phone",
    "contact_line",
    "start_date",
    "end_date",
    "start_date_long",
    "end_date_long",
    "event_type",
    "preparation_theme",
    "package_label",
    "cost_center",
    "date_du_jour",
)

_RULE_PATTERNS: list[tuple[str, str, str]] = [
    (r"\bZurich\b", "city", "replace"),
    (r"Titre de l[’']événement\s*:", "title", "replace"),
    (r"Date\s*:.*Lieu de l[’']événement\s*:", "date_range", "replace"),
    (r"Responsable\s*:", "responsible_person", "replace"),
    (r"AO Alliance|AOA—|Cours AOA", "title", "replace"),
    (r"Liste\s+(Enseignants|Participants)\s*\(", "weekday_date", "replace"),
    (r"\d{2}\s*[–-]\s*\d{2}\s+\w+", "date_range", "replace"),
    (r"\d{1,2}\s+(?:janvier|février|mars|avril|mai|juin|juillet|août|septembre|octobre|novembre|décembre)\s+\d{4}", "start_date_long", "replace"),
    (r"^(lundi|mardi|mercredi|jeudi|vendredi|samedi|dimanche)\b", "weekday_date", "replace"),
    (r"h[oô]tel|h[oô]pital", "lieu", "replace"),
    (r"pr[eé]nom\s*/?\s*nom", "participant_name", "keep"),
    (r"adresse@email.*téléphone|téléphone.*adresse@email", "contact_line", "replace"),
    (r"adresse@email", "responsible_email", "replace"),
    (r"téléphone du responsable national", "responsible_phone", "replace"),
    (r"à remplir par le/?\s*la\s+responsable?", "participant_name", "keep"),
    (r"\[PROJECT_NUMBER\]", "project_number", "replace"),
    (r"\[EVENT_TITLE\]", "title", "replace"),
    (r"\[CITY\]", "city", "replace"),
    (r"\[COUNTRY\]", "country", "replace"),
    (r"\[START_DATE\]", "start_date", "replace"),
    (r"\[END_DATE\]", "end_date", "replace"),
    (r"\[RESPONSIBLE\]", "responsible_person", "replace"),
    (r"\{\{\s*project_number\s*\}\}", "project_number", "replace"),
    (r"\{\{\s*title\s*\}\}", "title", "replace"),
    (r"\{\{\s*city\s*\}\}", "city", "replace"),
    (r"\{\{\s*Pays\s*\}\}", "country", "replace"),
    (r"\{\{\s*Ville\s*\}\}", "city", "replace"),
    (r"\{\{\s*Project Number\s*\}\}", "project_number", "replace"),
    (r"\{\{\s*Cost Center\s*\}\}", "cost_center", "replace"),
    (r"\{\{\s*Email\s*\}\}", "responsible_email", "replace"),
    (r"\{\{\s*Date du jour\s*\}\}", "date_du_jour", "replace"),
    (r"\{\{\s*Nom du responsable\s*\}\}", "responsible_person", "replace"),
    (r"\{\{\s*Numéro du responsable\s*\}\}", "responsible_phone", "replace"),
    (rf"\{{\{{\s*Nom de l{APOSTROPHE}{EVENEMENT}\s*\}}\}}", "title", "replace"),
    (rf"\{{\{{\s*Date de l{APOSTROPHE}{EVENEMENT}\s*\}}\}}", "date_du_jour", "replace"),
]


def _nvidia_url() -> str:
    explicit = os.getenv("NVIDIA_API_URL", "").strip().rstrip("/")
    if explicit:
        return explicit
    base = os.getenv("NVIDIA_API_BASE_URL", "https://integrate.api.nvidia.com/v1").strip().rstrip("/")
    return f"{base}/chat/completions"


def _collect_docx_paragraph_samples(root: ET.Element, samples: list[str], *, limit: int) -> None:
    for paragraph in root.iter(f"{{{W_NS}}}p"):
        texts: list[str] = []
        highlighted = False
        for run in paragraph.iter(f"{{{W_NS}}}r"):
            rpr = run.find(f"{{{W_NS}}}rPr")
            if rpr is not None and rpr.find(f"{{{W_NS}}}highlight") is not None:
                highlighted = True
            node = run.find(f"{{{W_NS}}}t")
            if node is not None and node.text:
                texts.append(node.text)
        line = "".join(texts).strip()
        if not line or len(line) < 2:
            continue
        is_event_line = any(
            token in line
            for token in ("AO Alliance", "AOA—", "AOA-", "Cours AOA", "Séminaire", "Titre de l")
        )
        if (
            highlighted
            or is_event_line
            or any(ch.isdigit() for ch in line)
            or "TBD" in line.upper()
        ):
            if line not in samples:
                samples.append(line[:240])
        if len(samples) >= limit:
            return


def _extract_docx_text_samples(data: bytes, *, limit: int = 12) -> list[str]:
    samples: list[str] = []
    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            xml_parts = sorted(
                n
                for n in zf.namelist()
                if n.startswith("word/") and n.endswith(".xml")
            )
            if not xml_parts:
                return samples
            for part in xml_parts:
                try:
                    root = ET.fromstring(zf.read(part))
                except ET.ParseError:
                    continue
                _collect_docx_paragraph_samples(root, samples, limit=limit)
                if len(samples) >= limit:
                    break
    except (zipfile.BadZipFile, ET.ParseError):
        return samples
    return samples


def _extract_xlsx_text_samples(data: bytes, *, limit: int = 10) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []

    samples: list[str] = []
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return samples

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(max_row=25, values_only=True):
            for cell in row:
                if not isinstance(cell, str) or len(cell.strip()) < 2:
                    continue
                text = cell.strip()
                if any(
                    token in text.upper()
                    for token in ("TBD", "ZURICH", "PROJECT", "[", "{{", "DATE", "HÔTEL", "HOTEL")
                ) or re.search(r"\d{2}\s+\w+\s+\d{4}", text):
                    if text not in samples:
                        samples.append(text[:240])
                if len(samples) >= limit:
                    return samples
    return samples


def _rules_analyze(samples: list[str], filename: str) -> list[dict[str, Any]]:
    fields: list[dict[str, Any]] = []
    seen: set[str] = set()
    for sample in samples:
        for pattern, context_key, strategy in _RULE_PATTERNS:
            if re.search(pattern, sample, re.I):
                key = f"{sample}:{context_key}"
                if key in seen:
                    continue
                seen.add(key)
                fields.append(
                    {
                        "sample": sample,
                        "context_key": context_key,
                        "strategy": strategy,
                        "classifier": "rules",
                    }
                )
                break
    if not fields and samples:
        fields.append(
            {
                "sample": samples[0],
                "context_key": "title",
                "strategy": "review",
                "classifier": "rules",
            }
        )
    return fields


async def _nvidia_analyze_fields(
    filename: str,
    document_role: str,
    samples: list[str],
) -> list[dict[str, Any]] | None:
    key = (os.getenv("NVIDIA_API_KEY") or "").strip()
    if not key or not samples:
        return None

    from tip_common.nvidia_client import nvidia_chat_completion

    samples_text = "\n".join(f"- {s}" for s in samples[:10])
    prompt = (
        "Tu analyses un modèle documentaire AO Alliance pour la génération automatique.\n"
        f"Fichier : {filename}\n"
        f"Rôle : {document_role}\n"
        f"Extraits du document :\n{samples_text}\n\n"
        f"Clés de contexte événement valides : {', '.join(CONTEXT_KEYS)}\n\n"
        "Pour chaque extrait à remplacer, indique la clé de contexte.\n"
        'Réponds UNIQUEMENT en JSON : {"fields":[{"sample":"…","context_key":"city",'
        '"strategy":"replace|keep|review","note":"…"}]}'
    )

    content, error = await nvidia_chat_completion(
        messages=[{"role": "user", "content": prompt}],
        temperature=0.15,
        max_tokens=600,
        timeout=45.0,
    )
    if not content:
        logger.warning("NVIDIA analyse template %s → règles : %s", filename, error)
        return None

    match = re.search(r"\{.*\}", content, re.DOTALL)
    if not match:
        return None
    try:
        parsed = json.loads(match.group())
    except json.JSONDecodeError:
        return None

    fields: list[dict[str, Any]] = []
    for item in parsed.get("fields") or []:
        if not isinstance(item, dict):
            continue
        sample = str(item.get("sample", "")).strip()
        context_key = str(item.get("context_key", "")).strip()
        if not sample or context_key not in CONTEXT_KEYS and context_key not in {
            "date_range",
            "weekday_date",
            "participant_name",
            "contact_placeholder",
            "lieu",
        }:
            continue
        fields.append(
            {
                "sample": sample,
                "context_key": context_key,
                "strategy": str(item.get("strategy", "replace")),
                "note": str(item.get("note", "")).strip(),
                "classifier": "nvidia",
            }
        )
    return fields or None


async def analyze_template_fields(
    *,
    filename: str,
    file_bytes: bytes,
    document_role: str = "autre",
    highlight_samples: list[str] | None = None,
) -> dict[str, Any]:
    """
    Analyse un fichier paquet et retourne les champs à remplacer automatiquement.
    Scan exhaustif (titres, en-têtes, corps) + NVIDIA si clé API.
    """
    from tip_common.document_section_scanner import extract_all_sections, sections_to_fields
    from tip_common.document_section_scanner import nvidia_classify_sections as nvidia_sections

    sections = extract_all_sections(filename, file_bytes, max_sections=80)
    samples = list(highlight_samples or [])
    for sec in sections:
        if sec.text not in samples:
            samples.append(sec.text)
    samples = samples[:30]

    ai_fields = await nvidia_sections(filename, document_role, sections)
    if not ai_fields:
        ai_fields = await _nvidia_analyze_fields(filename, document_role, samples)

    if ai_fields:
        classifier = "nvidia"
        fields = ai_fields
    else:
        classifier = "rules"
        fields = sections_to_fields(sections) or _rules_analyze(samples, filename)

    return {
        "filename": filename,
        "document_role": document_role,
        "classifier": classifier,
        "samples": samples,
        "section_count": len(sections),
        "fields": fields,
    }


def format_value_for_field(
    field: dict[str, Any],
    context: dict[str, Any],
) -> str | None:
    """Applique format_hint IA (titres, dates) si présent."""
    from tip_common.contact_fields import (
        format_contact_from_sample,
        format_phone_sample,
        format_responsible_sample,
    )

    key = str(field.get("context_key", "")).strip()
    hint = str(field.get("format_hint", "")).lower()
    sample = str(field.get("sample", "")).strip()
    kind = str(field.get("section_kind", "")).lower()

    if kind == "date_lieu_combined":
        from tip_common.location_fields import resolve_lieu_display

        date_val = (
            context.get("date_single_formatted")
            or context.get("start_date_long")
            or context.get("start_date")
            or ""
        )
        if isinstance(date_val, str):
            date_val = date_val.strip()
        lieu = resolve_lieu_display(context)
        if not date_val or not lieu or not sample:
            return None
        updated = re.sub(
            r"\d{1,2}\s+(?:janvier|février|mars|avril|mai|juin|juillet|août|septembre|octobre|novembre|décembre)\s+\d{4}",
            date_val,
            sample,
            count=1,
            flags=re.I,
        )
        for old_lieu in (
            "Bangui, RCA",
            "Dakar, Sénégal",
            "Mbour, Sénégal",
            "Kaffrine, Sénégal",
            "Brazzaville, Congo",
            "Zurich, Suisse",
            "Ethiopia, Ethiopia",
            "Addis Ababa, Ethiopia",
        ):
            if old_lieu in updated:
                updated = updated.replace(old_lieu, lieu, 1)
                break
        if updated == sample:
            return None
        if len(updated) > len(sample):
            return None
        if len(updated) < len(sample):
            updated = updated + " " * (len(sample) - len(updated))
        return updated

    if key in {"contact_line", "contact_placeholder"} or (
        sample and ("adresse@email" in sample.lower() or ("courriel" in sample.lower() and "téléphone" in sample.lower()))
    ):
        contact = format_contact_from_sample(sample, context)
        if contact:
            return contact

    if key in {"responsible_phone", "responsible_email"} or (
        sample and re.search(r"téléphone du responsable", sample, re.I)
    ):
        phone_line = format_phone_sample(sample, context)
        if phone_line:
            return phone_line

    if key == "responsible_email" and sample.lower() == "adresse@email":
        email = (context.get("responsible_email") or "").strip()
        if email:
            return email

    if key == "responsible_person":
        responsible = format_responsible_sample(sample, context)
        if responsible:
            return responsible

    base = context_value_for_key(key, context)
    if not base:
        return None
    if key == "title" or "titre" in hint:
        title = (context.get("title_formatted") or context.get("title") or "").strip()
        if title and re.match(r"^Titre de l[’']événement\s*:\s*", sample, re.I):
            return re.sub(
                r"^(Titre de l[’']événement\s*:\s*).+$",
                rf"\g<1>{title}",
                sample.strip(),
                flags=re.I,
            )
        if title:
            return title
    if "ville" in hint and "pays" in hint:
        from tip_common.location_fields import resolve_lieu_doc_display

        lieu = resolve_lieu_doc_display(context)
        if lieu:
            return lieu
        city = (context.get("city") or "").strip()
        country = (context.get("country") or "").strip()
        if city and country:
            return f"{city}, {country}"
    if key == "lieu" and ("date" in hint or "mois" in hint or "jj" in hint):
        header = (context.get("header_lieu_date") or "").strip()
        if header:
            return header
    if key == "date_range":
        formatted = (context.get("date_range_formatted") or "").strip()
        if formatted:
            return formatted
    if key == "weekday_date":
        single = (context.get("date_single_formatted") or "").strip()
        if single:
            return single
    if "jj" in hint or "mois" in hint or key in {"date_range", "start_date_long", "weekday_date"}:
        return base
    return base


def context_value_for_key(key: str, context: dict[str, Any]) -> str | None:
    """Valeur de remplacement pour une clé de contexte (analyse IA ou règles)."""
    key = (key or "").strip()
    if not key:
        return None
    direct = {
        "project_number": context.get("project_number"),
        "title": context.get("title_formatted") or context.get("title"),
        "city": context.get("city"),
        "country": context.get("country"),
        "lieu": context.get("lieu_formatted") or context.get("lieu") or context.get("location"),
        "region": context.get("region"),
        "responsible_person": context.get("responsible_person") or context.get("responsable"),
        "responsible_email": context.get("responsible_email"),
        "responsible_phone": context.get("responsible_phone"),
        "participants_expected": context.get("participants_expected"),
        "prepared_by": context.get("prepared_by") or context.get("prepared_by_name"),
        "contact_line": context.get("contact_line"),
        "start_date": context.get("start_date"),
        "end_date": context.get("end_date"),
        "start_date_long": context.get("start_date_long"),
        "end_date_long": context.get("end_date_long"),
        "event_type": context.get("event_type"),
        "preparation_theme": context.get("preparation_theme") or context.get("theme"),
        "package_label": context.get("package_label"),
        "cost_center": context.get("cost_center"),
        "date_du_jour": context.get("date_du_jour") or context.get("today"),
        "date_single_formatted": context.get("date_single_formatted"),
        "date_range_formatted": context.get("date_range_formatted"),
        "header_lieu_date": context.get("header_lieu_date"),
        "organizer_responsible_name": context.get("organizer_responsible_name"),
        "year": context.get("year") or context.get("annee"),
    }
    if key in direct and direct[key]:
        return str(direct[key]).strip() or None
    if key.startswith("teacher_") or key.startswith("enseignant_"):
        try:
            index = int(key.rsplit("_", 1)[-1]) - 1
        except ValueError:
            return None
        names = list(context.get("teacher_names") or [])
        if not names:
            for teacher in context.get("teachers") or []:
                if isinstance(teacher, dict):
                    name = f"{teacher.get('first_name', '')} {teacher.get('last_name', '')}".strip()
                    if name:
                        names.append(name)
        if 0 <= index < len(names):
            return str(names[index]).strip() or None
        return None
    if key == "participant_name":
        return None
    if key == "contact_line":
        email = (context.get("responsible_email") or "").strip()
        phone = (context.get("responsible_phone") or "").strip()
        if email and phone:
            return f"Courriel: {email}        Téléphone: {phone}"
        if email:
            return f"Courriel: {email}"
        if phone:
            return f"Téléphone: {phone}"
        return None
    if key == "date_range":
        formatted = context.get("date_range_formatted")
        if formatted:
            return str(formatted).strip()
        start = context.get("start_date_long") or context.get("start_date")
        end = context.get("end_date_long") or context.get("end_date") or start
        lieu = context.get("lieu_formatted") or context.get("lieu") or context.get("location") or ""
        if start == end or not end or str(end) == str(start):
            return str(start or "").strip() or None
        if start and lieu:
            return f"{start} — {end}          {lieu}".strip()
        return str(start or end or "").strip() or None
    if key == "weekday_date":
        return str(
            context.get("date_single_formatted")
            or context.get("start_date_long")
            or context.get("start_date")
            or ""
        ).strip() or None
    if key == "start_date_long":
        return str(
            context.get("date_single_formatted") or context.get("start_date_long") or ""
        ).strip() or None
    if key == "lieu":
        from tip_common.location_fields import resolve_lieu_doc_display

        formatted = resolve_lieu_doc_display(context)
        if formatted:
            return formatted
        formatted = context.get("lieu_formatted") or context.get("lieu") or context.get("location")
        if formatted:
            return str(formatted).strip()
        city = (context.get("city") or "").strip()
        country = (context.get("country") or "").strip()
        if city or country:
            return ", ".join(part for part in (city, country) if part)
    return None


def _is_static_label_sample(sample: str) -> bool:
    """Libellés / textes fixes qui ne doivent jamais être remplacés."""
    text = sample.strip()
    if not text:
        return True
    if text.endswith(":") and len(text) < 80:
        return True
    if re.match(r"^à remplir par le/?\s*la\s+responsable?$", text, re.I):
        return True
    if re.match(r"^Formulaire de budget", text, re.I):
        return True
    if text.startswith("Note importante:"):
        return True
    if re.match(r"^Nombre de ", text, re.I):
        return True
    if re.match(r"^Budget maximum", text, re.I):
        return True
    if re.match(
        r"^(Monnaie:|Frais et perdiems|Hôtel|Restauration|Transports|Traduction|Lieu de l|"
        r"Matériel|Infrastructure|\(moins\)|Estimation Budget|Per diem|En signant|"
        r"Coordonnées bancaires|Rapport du responsable|Rapport du responsable national de l|Veuillez |Merci de retourner|"
        r"Voulez-vous recommander|Prochaine fois|AO Alliance Foundation|"
        r"Enseignants |Nom du titulaire|Adresse |Nom de la banque|N° de compte|SWIFT|Clearing|"
        r"Appel à une banque|Pré cours|Conférences|Discussions|Prochaine fois|Rapport rédigé|"
        r"Lieu, date|Voulez-vous recommander|Support technique|Problèmes avec|Traduction:|"
        r"Infrastructure / repas|Nom de l.?événement|Date de l.?événement)",
        text,
        re.I,
    ):
        return True
    if re.match(r"^\d+\s*$", text):
        return True
    return False


_AO_TITLE_MARKERS = ("AO Alliance", "AOA—", "AOA-", "Cours AOA", "Séminaire AO", "Séminaire AOA")

_SAFE_EXHAUSTIVE_KINDS = frozenset(
    {
        "highlight",
        "lieu_line",
        "date_line",
        "date_lieu_combined",
        "placeholder",
        "event_header",
        "document_title",
        "label_field",
    }
)


def _field_safe_for_exhaustive(field: dict[str, Any]) -> bool:
    if str(field.get("strategy", "replace")).lower() == "keep":
        return False
    kind = str(field.get("section_kind", "")).lower()
    if kind in {"body_text", "static", "participant_row", "table_label"}:
        return False
    key = str(field.get("context_key", "")).strip()
    sample = str(field.get("sample", "")).strip()
    if not sample:
        return False
    if key == "title" and kind not in _SAFE_EXHAUSTIVE_KINDS:
        return False
    if key == "title" and len(sample) > 100 and kind != "highlight":
        return False
    if key == "title":
        if not any(marker in sample for marker in _AO_TITLE_MARKERS):
            return False
        if len(sample) < 20:
            return False
    if key == "title" and kind in {"document_title", "event_header"}:
        return any(marker in sample for marker in _AO_TITLE_MARKERS)
    if kind in _SAFE_EXHAUSTIVE_KINDS:
        return True
    return False


def build_replacement_pairs(
    fields: list[dict[str, Any]] | None,
    context: dict[str, Any],
) -> list[tuple[str, str]]:
    """Paires (texte modèle → valeur événement) issues de l'analyse IA par fichier."""
    pairs: list[tuple[str, str]] = []
    seen: set[str] = set()
    for field in fields or []:
        if not _field_safe_for_exhaustive(field):
            continue
        sample = str(field.get("sample", "")).strip()
        if not sample or sample in seen or _is_static_label_sample(sample):
            continue
        from tip_common.french_placeholders import is_french_brace_placeholder

        key = str(field.get("context_key", "")).strip()
        if is_french_brace_placeholder(sample) and key in {"city", "country", "lieu", "region"}:
            continue
        value = format_value_for_field(field, context) or context_value_for_key(
            str(field.get("context_key", "")), context
        )
        if not value or sample == value:
            continue
        seen.add(sample)
        pairs.append((sample, value))
    return pairs
