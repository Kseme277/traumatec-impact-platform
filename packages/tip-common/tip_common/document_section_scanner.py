"""Extraction exhaustive des sections remplaçables dans les modèles AO Alliance."""

from __future__ import annotations

import json
import logging
import os
import re
import zipfile
from dataclasses import asdict, dataclass, field
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from xml.etree import ElementTree as ET

logger = logging.getLogger(__name__)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

DATE_RE = re.compile(
    r"\d{1,2}\s*(?:[–\-]\s*\d{1,2}\s+)?"
    r"(?:janvier|février|mars|avril|mai|juin|juillet|août|septembre|octobre|novembre|décembre|ocobre)"
    r"\s+\d{4}",
    re.I,
)
SKIP_LINE_RE = re.compile(
    r"^(\d{1,2}|[ivxlc]+\.?|no\.?|noms?\s+et\s+prénoms?|pays|h[oô]pital|email|signature|titre\s*\(|fonction)$",
    re.I,
)
EVENT_MARKERS = (
    "AO Alliance", "AOA—", "AOA-", "Cours AOA", "Séminaire", "Sminaire",
    "événement", "evenement", "Bangui", "Dakar", "Mbour", "Zurich", "TBD",
    "Zurich", "Sénégal", "Senegal", "Congo", "RCA", "Kaffrine", "Brazzaville",
)
STATIC_KEEP = (
    "AO Alliance Foundation",
    "Fiduciar Treuhand",
    "Theaterweg",
    "Chur",
    "c/o Fiduciar",
    "Barbara Rigassi",
    "Présidente",
)

SECTION_KINDS = (
    "document_title",
    "event_header",
    "date_line",
    "lieu_line",
    "date_lieu_combined",
    "highlight",
    "label_field",
    "body_text",
    "table_label",
    "placeholder",
    "participant_row",
    "static",
)


@dataclass
class DocumentSection:
    text: str
    location: str
    section_kind: str = "body_text"
    highlighted: bool = False
    context_key: str = ""
    strategy: str = "review"
    note: str = ""

    def to_field(self) -> dict[str, Any]:
        return {
            "sample": self.text,
            "location": self.location,
            "section_kind": self.section_kind,
            "context_key": self.context_key,
            "strategy": self.strategy,
            "highlighted": self.highlighted,
            "note": self.note,
        }


def _is_noise(text: str) -> bool:
    t = text.strip()
    if len(t) < 4:
        return True
    if SKIP_LINE_RE.match(t):
        return True
    if re.fullmatch(r"[\d\.\s]+", t):
        return True
    return False


def _guess_section_kind(text: str, *, highlighted: bool, location: str) -> tuple[str, str, str]:
    """Retourne (section_kind, context_key, strategy)."""
    t = text.strip()

    if re.match(r"^à remplir par le/?\s*la\s+responsable?$", t, re.I):
        return "placeholder", "", "keep"
    if re.match(
        r"^(Formulaire de budget|Note importante|Budget maximum|Monnaie:|Frais et perdiems|"
        r"Nombre de |Quantité \(|Coût unitaire|Per diem|Hôtel pour|Repas|Pause-café|"
        r"Transports locaux|Traduction simultanée|Salles|Equipement audiovisuel|Matériel de|"
        r"Infrastructure de|Impression du|\(moins\) Revenus|Estimation Budget|En signant|"
        r"Budget préparé par|Coordonnées bancaires pour|Rapport du responsable national|"
        r"Rapport du responsable national de l|Veuillez remplir|Veuillez reporter|Veuillez donner|Veuillez svp|Merci de retourner|"
        r"Voulez-vous recommander|Prochaine fois|AO Alliance Foundation|"
        r"Nom du titulaire|Adresse complète|Nom de la banque|Adresse de la banque|"
        r"N° de compte|SWIFT|Clearing|Appel à une banque|Enseignants chirurgiens|"
        r"Enseignants ORP|Enseignants régionaux|Voulez-vous recommander|Pré cours|"
        r"Conférences:|Responsable des travaux|Instructeurs de table|Discussions:|"
        r"Support technique|Prochaine fois|Rapport rédigé par|Lieu, date:|"
        r"Bienvenue|Veuillez agréer|But du cours|Audience cible|Les grands chapitres|"
        r"Objectifs du cours|Collège d|Informations générales|Propriété intellectuelle|"
        r"Organisation du cours|Lieu du cours|Bureau d|Personne de contact|"
        r"Les principes AO|Module \d|TEMPS|SUJETS|QUI|PAUSE)",
        t,
        re.I,
    ):
        return "static", "", "keep"
    if re.match(
        r"^(Nom de l.?événement|Responsable national de l.?événement|Titre de l.?événement|"
        r"Date de l.?événement|Lieu de l.?événement)",
        t,
        re.I,
    ):
        return "label_field", "", "keep"
    if t.endswith(":") and len(t) < 80 and "événement" in t.lower():
        return "label_field", "", "keep"

    if re.search(r"pr[eé]nom\s*/?\s*nom", t, re.I) or t in {"Prénom Nom", "Prénom/ Nom", "Fonction"}:
        return "participant_row", "participant_name", "keep"
    if "adresse@email" in t.lower() and ("téléphone" in t.lower() or "telephone" in t.lower()):
        return "placeholder", "contact_line", "replace"
    if "adresse@email" in t.lower():
        return "placeholder", "responsible_email", "replace"
    if re.search(r"téléphone du responsable national", t, re.I):
        return "label_field", "responsible_phone", "replace"
    if highlighted:
        if DATE_RE.search(t) and any(x in t.lower() for x in ("hôtel", "hotel", "hôpital", "hopital")):
            return "highlight", "lieu", "replace"
        if DATE_RE.search(t):
            return "highlight", "date_range", "replace"
        if re.match(r"^(lundi|mardi|mercredi|jeudi|vendredi|samedi|dimanche)\b", t, re.I):
            return "highlight", "weekday_date", "replace"
        if re.search(r"h[oô]tel|h[oô]pital", t, re.I):
            return "highlight", "lieu", "replace"
        return "highlight", "", "keep"

    if re.match(r"^(Titre de l[’']événement\s*:)", t, re.I):
        return "label_field", "", "keep"
    if re.match(r"^Date\s*:", t, re.I) and "Lieu" in t:
        return "date_lieu_combined", "date_range", "replace"
    if re.match(r"^Date\s*:", t, re.I):
        return "date_line", "date_range", "replace"
    if re.match(r"^(Responsable national|Responsable|Nombre de participants)\s*:", t, re.I):
        return "label_field", "", "keep"
    if re.match(r"^Liste\s+(Enseignants|Participants)\s*\(", t, re.I):
        return "event_header", "weekday_date", "replace"
    if (
        any(m in t for m in ("AO Alliance", "AOA—", "AOA-", "Cours AOA", "Séminaire AO"))
        and len(t) > 35
        and "budget prévisionnel!" in location
        and re.search(r"!R[4567]C3$", location)
    ):
        return "event_header", "title", "replace"
    if any(m in t for m in ("AO Alliance", "AOA—", "AOA-", "Cours AOA", "Séminaire AO")) and len(t) > 35:
        kind = "document_title" if "header" in location or "document" in location else "event_header"
        return kind, "title", "replace"
    if DATE_RE.search(t) and "," in t and len(t) < 120:
        return "lieu_line", "lieu", "replace"
    if DATE_RE.fullmatch(t) or (DATE_RE.search(t) and len(t) < 40):
        return "date_line", "start_date_long", "replace"
    if re.search(r"pr[eé]nom\s*/?\s*nom", t, re.I):
        return "participant_row", "participant_name", "keep"
    if "adresse@email" in t.lower() and ("téléphone" in t.lower() or "telephone" in t.lower()):
        return "placeholder", "contact_line", "replace"
    if "adresse@email" in t.lower():
        return "placeholder", "responsible_email", "replace"
    if re.search(r"téléphone du responsable national", t, re.I):
        return "label_field", "responsible_phone", "replace"
    if t.startswith("Note importante"):
        return "static", "", "keep"
    if t in {"TBD", "Zurich"} or t.startswith("[") or "{{" in t:
        return "placeholder", "city", "replace"
    if "header" in location and len(t) > 20:
        return "event_header", "title", "replace"
    if any(m in t for m in EVENT_MARKERS) and len(t) > 15:
        if any(s in t for s in STATIC_KEEP) and "Séminaire" not in t and "Cours" not in t:
            return "static", "", "keep"
        return "body_text", "", "keep"
    if DATE_RE.search(t):
        return "date_line", "date_range", "replace"
    if re.search(r"\b(RCA|Sénégal|Senegal|Congo|Zurich|Bangui|Dakar|Mbour|Gambia)\b", t, re.I):
        return "lieu_line", "lieu", "replace"
    return "body_text", "", "keep"


def _paragraph_sections(
    root: ET.Element,
    part_name: str,
    *,
    max_sections: int = 80,
) -> list[DocumentSection]:
    sections: list[DocumentSection] = []
    seen: set[str] = set()

    for idx, paragraph in enumerate(root.iter(f"{{{W_NS}}}p")):
        texts: list[str] = []
        highlighted = False
        for run in paragraph.iter(f"{{{W_NS}}}r"):
            rpr = run.find(f"{{{W_NS}}}rPr")
            if rpr is not None and rpr.find(f"{{{W_NS}}}highlight") is not None:
                highlighted = True
            node = run.find(f"{{{W_NS}}}t")
            if node is not None and node.text:
                texts.append(node.text)
        line = "".join(texts).strip()
        if _is_noise(line) or line in seen:
            continue
        seen.add(line)
        kind, ctx_key, strategy = _guess_section_kind(line, highlighted=highlighted, location=part_name)
        if strategy == "keep" and kind in {"body_text", "static", "participant_row"}:
            continue
        sections.append(
            DocumentSection(
                text=line[:400],
                location=f"{part_name}#p{idx}",
                section_kind=kind,
                highlighted=highlighted,
                context_key=ctx_key,
                strategy=strategy,
            )
        )
        if len(sections) >= max_sections:
            break
    return sections


def extract_docx_sections(data: bytes, *, max_sections: int = 150) -> list[DocumentSection]:
    sections: list[DocumentSection] = []
    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            xml_parts = sorted(
                n for n in zf.namelist() if n.startswith("word/") and n.endswith(".xml")
            )
            for part in xml_parts:
                try:
                    root = ET.fromstring(zf.read(part))
                except ET.ParseError:
                    continue
                sections.extend(_paragraph_sections(root, part, max_sections=max_sections))
                if len(sections) >= max_sections:
                    break
    except zipfile.BadZipFile:
        pass
    return sections[:max_sections]


def extract_xlsx_sections(data: bytes, *, max_rows: int = 120, max_sections: int = 120) -> list[DocumentSection]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []

    sections: list[DocumentSection] = []
    seen: set[str] = set()
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return []

    for sheet in wb.worksheets:
        for row_idx, row in enumerate(sheet.iter_rows(max_row=max_rows, values_only=True), start=1):
            for col_idx, cell in enumerate(row, start=1):
                if not isinstance(cell, str):
                    continue
                text = cell.strip()
                if _is_noise(text) or text in seen or len(text) < 4:
                    continue
                if any(s in text for s in STATIC_KEEP) and not any(
                    m in text for m in ("Séminaire", "Cours ", "événement", "Bangui", "Dakar")
                ):
                    continue
                if not (
                    DATE_RE.search(text)
                    or any(tok in text.upper() for tok in ("TBD", "ZURICH", "AO ALLIANCE", "AOA", "PROJECT", "ÉVÉNEMENT", "EVENEMENT", "RESPONSABLE", "LIEU", "DATE", "BUDGET", "ACCORD", "[", "{{", "RCA", "CONGO", "SÉNÉGAL"))
                    or any(m in text for m in EVENT_MARKERS)
                    or (len(text) > 20 and re.search(r"[A-Za-zÀ-ÿ]{4}", text))
                ):
                    continue
                seen.add(text)
                kind, ctx_key, strategy = _guess_section_kind(
                    text, highlighted=False, location=f"{sheet.title}!{row_idx}:{col_idx}"
                )
                sections.append(
                    DocumentSection(
                        text=text[:400],
                        location=f"{sheet.title}!R{row_idx}C{col_idx}",
                        section_kind=kind,
                        context_key=ctx_key,
                        strategy=strategy,
                    )
                )
                if len(sections) >= max_sections:
                    return sections
    return sections


def extract_doc_strings(data: bytes, *, min_len: int = 8) -> list[DocumentSection]:
    """Extrait chaînes UTF-16/UTF-8 lisibles des .doc binaires."""
    sections: list[DocumentSection] = []
    seen: set[str] = set()

    def add_text(raw: str, loc: str) -> None:
        t = raw.strip()
        if len(t) < min_len or _is_noise(t) or t in seen:
            return
        if not re.search(r"[a-zA-ZÀ-ÿ]{4}", t):
            return
        seen.add(t)
        kind, ctx_key, strategy = _guess_section_kind(t, highlighted=False, location=loc)
        if strategy == "keep" and kind == "body_text":
            return
        sections.append(
            DocumentSection(
                text=t[:400],
                location=loc,
                section_kind=kind,
                context_key=ctx_key,
                strategy=strategy,
            )
        )

    for encoding, label in (("utf-16-le", "utf16"), ("utf-8", "utf8")):
        try:
            decoded = data.decode(encoding, errors="ignore")
        except Exception:
            continue
        for match in re.finditer(r"[\wÀ-ÿ][\wÀ-ÿ\s,–\-':;/\(\)\.]{12,}", decoded):
            add_text(match.group(), f"binary:{label}")

    for token in ("TBD", "Zurich", "[PROJECT_NUMBER]", "[EVENT_TITLE]", "[CITY]", "adresse@email"):
        if token.encode() in data or token.lower().encode() in data.lower():
            add_text(token, "binary:token")

    return sections[:80]


def extract_all_sections(
    filename: str,
    file_bytes: bytes,
    *,
    max_sections: int = 150,
) -> list[DocumentSection]:
    ext = PurePosixPath(filename).suffix.lower()
    if ext == ".docx":
        return extract_docx_sections(file_bytes, max_sections=max_sections)
    if ext == ".xlsx":
        return extract_xlsx_sections(file_bytes, max_sections=max_sections)
    if ext == ".doc":
        return extract_doc_strings(file_bytes)
    return []


def sections_to_fields(sections: list[DocumentSection]) -> list[dict[str, Any]]:
    return [s.to_field() for s in sections if s.strategy == "replace" or s.highlighted]


def _merge_field_lists(*lists: list[dict[str, Any]] | None) -> list[dict[str, Any]]:
    merged: list[dict[str, Any]] = []
    seen: set[str] = set()
    for lst in lists:
        for field in lst or []:
            sample = str(field.get("sample", "")).strip()
            if not sample or sample in seen:
                continue
            if str(field.get("strategy", "replace")).lower() == "keep":
                continue
            seen.add(sample)
            merged.append(field)
    return merged


async def _mistral_classify_batch(
    filename: str,
    document_role: str,
    sections: list[DocumentSection],
    *,
    batch_index: int,
) -> list[dict[str, Any]] | None:
    from tip_common.mistral_client import mistral_api_key, mistral_chat_completion
    from tip_common.template_field_analyzer import CONTEXT_KEYS

    if not mistral_api_key() or not sections:
        return None

    lines = [f"[{s.location}] {s.text[:220]}" for s in sections]
    prompt = (
        "Analyse EXHAUSTIVE d'un modèle AO Alliance. Liste TOUTES les sections à remplacer par les données événement.\n"
        f"Fichier : {filename} | Rôle : {document_role} | Lot {batch_index + 1}\n\n"
        + "\n".join(f"- {ln}" for ln in lines)
        + "\n\n"
        f"Clés : {', '.join(CONTEXT_KEYS)}, date_range, weekday_date, lieu, start_date_long\n\n"
        "RÈGLES :\n"
        "- strategy=replace pour : titres, dates, lieux, villes, pays, responsable, numéro projet, hôtel\n"
        "- strategy=replace pour courriel/téléphone responsable (adresse@email, N° téléphone…)\n"
        "- strategy=keep pour : Prénom Nom (tableaux vides), adresse siège AO Foundation Chur, texte légal générique\n"
        "- format_hint : format cible (ex: « Titre AO complet », « Ville, Pays JJ mois AAAA »)\n"
        "- sample : texte EXACT du modèle\n\n"
        'JSON : {"sections":[{"sample":"…","context_key":"title","strategy":"replace","format_hint":"…"}]}'
    )

    content, error = await mistral_chat_completion(
        messages=[{"role": "user", "content": prompt}],
        temperature=0.05,
        max_tokens=3500,
        timeout=120.0,
    )
    if not content:
        logger.warning("Mistral batch %s %s : %s", filename, batch_index, error)
        return None

    match = re.search(r"\{.*\}", content, re.DOTALL)
    if not match:
        return None
    try:
        parsed = json.loads(match.group())
    except json.JSONDecodeError:
        return None

    valid_keys = set(CONTEXT_KEYS) | {
        "date_range",
        "weekday_date",
        "participant_name",
        "contact_line",
        "contact_placeholder",
        "lieu",
    }
    out: list[dict[str, Any]] = []
    for item in parsed.get("sections") or parsed.get("fields") or []:
        if not isinstance(item, dict):
            continue
        sample = str(item.get("sample", "")).strip()
        ctx = str(item.get("context_key", "")).strip()
        if not sample or str(item.get("strategy", "replace")).lower() == "keep":
            continue
        out.append(
            {
                "sample": sample,
                "context_key": ctx if ctx in valid_keys else "title",
                "strategy": "replace",
                "section_kind": str(item.get("section_kind", "body_text")),
                "format_hint": str(item.get("format_hint", "")).strip(),
                "note": str(item.get("note", "")).strip(),
                "classifier": "mistral",
            }
        )
    return out or None


async def mistral_classify_sections(
    filename: str,
    document_role: str,
    sections: list[DocumentSection],
) -> list[dict[str, Any]] | None:
    if not sections:
        return None
    batch_size = 30
    all_fields: list[dict[str, Any]] = []
    for i in range(0, len(sections), batch_size):
        batch = sections[i : i + batch_size]
        result = await _mistral_classify_batch(filename, document_role, batch, batch_index=i // batch_size)
        if result:
            all_fields.extend(result)
    return all_fields or None


async def _nvidia_classify_batch(
    filename: str,
    document_role: str,
    sections: list[DocumentSection],
    *,
    batch_index: int,
) -> list[dict[str, Any]] | None:
    key = (os.getenv("NVIDIA_API_KEY") or "").strip()
    if not key or not sections:
        return None

    from tip_common.nvidia_client import nvidia_chat_completion
    from tip_common.template_field_analyzer import CONTEXT_KEYS
    lines = [f"[{s.location}] {s.text[:220]}" for s in sections]
    prompt = (
        "Analyse EXHAUSTIVE d'un modèle AO Alliance. Liste TOUTES les sections à remplacer par les données événement.\n"
        f"Fichier : {filename} | Rôle : {document_role} | Lot {batch_index + 1}\n\n"
        + "\n".join(f"- {ln}" for ln in lines)
        + "\n\n"
        f"Clés : {', '.join(CONTEXT_KEYS)}, date_range, weekday_date, lieu, start_date_long\n\n"
        "RÈGLES :\n"
        "- strategy=replace pour : titres, dates, lieux, villes, pays, responsable, numéro projet, hôtel\n"
        "- strategy=replace pour courriel/téléphone responsable (adresse@email, N° téléphone…)\n"
        "- strategy=keep pour : Prénom Nom (tableaux vides), adresse siège AO Foundation Chur, texte légal générique\n"
        "- format_hint : format cible (ex: « Titre AO complet », « Ville, Pays JJ mois AAAA »)\n"
        "- sample : texte EXACT du modèle\n\n"
        'JSON : {"sections":[{"sample":"…","context_key":"title","strategy":"replace","format_hint":"…"}]}'
    )

    content, error = await nvidia_chat_completion(
        messages=[{"role": "user", "content": prompt}],
        temperature=0.05,
        max_tokens=3500,
        timeout=120.0,
    )
    if not content:
        logger.warning("NVIDIA batch %s %s : %s", filename, batch_index, error)
        return None

    match = re.search(r"\{.*\}", content, re.DOTALL)
    if not match:
        return None
    try:
        parsed = json.loads(match.group())
    except json.JSONDecodeError:
        return None

    valid_keys = set(CONTEXT_KEYS) | {
        "date_range",
        "weekday_date",
        "participant_name",
        "contact_line",
        "contact_placeholder",
        "lieu",
    }
    out: list[dict[str, Any]] = []
    for item in parsed.get("sections") or parsed.get("fields") or []:
        if not isinstance(item, dict):
            continue
        sample = str(item.get("sample", "")).strip()
        ctx = str(item.get("context_key", "")).strip()
        if not sample or str(item.get("strategy", "replace")).lower() == "keep":
            continue
        out.append(
            {
                "sample": sample,
                "context_key": ctx if ctx in valid_keys else "title",
                "strategy": "replace",
                "section_kind": str(item.get("section_kind", "body_text")),
                "format_hint": str(item.get("format_hint", "")).strip(),
                "note": str(item.get("note", "")).strip(),
                "classifier": "nvidia",
            }
        )
    return out or None


async def nvidia_classify_sections(
    filename: str,
    document_role: str,
    sections: list[DocumentSection],
) -> list[dict[str, Any]] | None:
    if not sections:
        return None
    batch_size = 30
    all_fields: list[dict[str, Any]] = []
    for i in range(0, len(sections), batch_size):
        batch = sections[i : i + batch_size]
        result = await _nvidia_classify_batch(filename, document_role, batch, batch_index=i // batch_size)
        if result:
            all_fields.extend(result)
    return all_fields or None


async def scan_document_sections(
    *,
    filename: str,
    file_bytes: bytes,
    document_role: str = "autre",
    use_ai: bool = False,
) -> dict[str, Any]:
    from tip_common.brace_placeholder_scanner import scan_brace_placeholders

    brace_scan = scan_brace_placeholders(
        filename=filename,
        file_bytes=file_bytes,
        document_role=document_role,
    )
    brace_fields = brace_scan.get("replacement_fields") or []

    sections = extract_all_sections(filename, file_bytes)
    rule_fields = sections_to_fields(sections)

    ai_fields: list[dict[str, Any]] | None = None
    if use_ai:
        ai_fields = await mistral_classify_sections(filename, document_role, sections)

    fields = _merge_field_lists(brace_fields, rule_fields, ai_fields)

    classifier_parts: list[str] = []
    if brace_fields:
        classifier_parts.append("brace")
    if rule_fields:
        classifier_parts.append("rules")
    if ai_fields:
        classifier_parts.append("mistral")
    classifier = "+".join(classifier_parts) if classifier_parts else "none"

    return {
        "filename": filename,
        "document_role": document_role,
        "classifier": classifier,
        "section_count": len(sections) or brace_scan.get("section_count", 0),
        "replaceable_count": len([f for f in fields if f.get("strategy") != "keep"]),
        "sections": [asdict(s) for s in sections],
        "replacement_fields": fields,
    }
