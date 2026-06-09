"""Remplacements cellules Excel (.xlsx) pour les modèles AO Alliance."""

from __future__ import annotations

import logging
import re
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


def _replacement_pairs(context: dict[str, Any]) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    mapping = {
        "TBD": context.get("city") or context.get("lieu") or context.get("project_number", ""),
        "Zurich": context.get("city") or "Zurich",
        "[PROJECT_NUMBER]": context.get("project_number", ""),
        "[EVENT_TITLE]": context.get("title", ""),
        "[CITY]": context.get("city", ""),
        "[COUNTRY]": context.get("country", ""),
        "[START_DATE]": context.get("start_date", ""),
        "[END_DATE]": context.get("end_date", ""),
        "[RESPONSIBLE]": context.get("responsible_person", ""),
        "{{ project_number }}": context.get("project_number", ""),
        "{{ title }}": context.get("title", ""),
        "{{ city }}": context.get("city", ""),
        "{{ country }}": context.get("country", ""),
        "{{ start_date }}": context.get("start_date", ""),
        "{{ end_date }}": context.get("end_date", ""),
    }
    for old, new in mapping.items():
        if new and str(old) != str(new):
            pairs.append((old, str(new)))
    return pairs


def _replace_cell_value(
    value: Any,
    pairs: list[tuple[str, str]],
    *,
    context: dict[str, Any],
    document_role: str | None = None,
    replacement_fields: list[dict[str, Any]] | None = None,
) -> Any:
    if not isinstance(value, str) or not value.strip():
        return value
    from app.services.docgen.document_role_replace import excel_cell_replace

    updated = value
    for old, new in pairs:
        if old in updated:
            updated = updated.replace(old, new)
    return excel_cell_replace(
        updated,
        context,
        document_role=document_role,
        replacement_fields=replacement_fields,
    )


def apply_excel_replacements(
    xlsx_bytes: bytes,
    context: dict[str, Any],
    *,
    replacement_fields: list[dict[str, Any]] | None = None,
    document_role: str | None = None,
) -> bytes:
    """Remplace les placeholders texte dans toutes les cellules."""
    if document_role == "budget":
        from app.services.docgen.budget_excel_replace import apply_budget_excel_replacements

        return apply_budget_excel_replacements(xlsx_bytes, context)

    if document_role == "accord_collaboration":
        from app.services.docgen.accord_excel_replace import apply_accord_excel_replacements

        return apply_accord_excel_replacements(xlsx_bytes, context)

    if document_role == "rapport_depenses":
        from app.services.docgen.strict_excel_replace import apply_strict_excel_replacements

        return apply_strict_excel_replacements(
            xlsx_bytes,
            context,
            replacement_fields=replacement_fields,
            document_role=document_role,
        )

    from tip_common.template_field_analyzer import build_replacement_pairs

    pairs = _replacement_pairs(context)
    for old, new in build_replacement_pairs(replacement_fields, context):
        if (old, new) not in pairs:
            pairs.append((old, new))
    if not pairs:
        return xlsx_bytes

    try:
        wb = load_workbook(BytesIO(xlsx_bytes))
    except Exception as exc:
        logger.warning("Excel illisible (%s) — copie brute", exc)
        return xlsx_bytes

    replaced = 0
    for sheet in wb.worksheets:
        if re.search(r"jour\s*[23]", sheet.title, re.I) and context.get("package_duration_days") == 1:
            continue
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                new_val = _replace_cell_value(
                    cell.value,
                    pairs,
                    context=context,
                    document_role=document_role,
                    replacement_fields=replacement_fields,
                )
                if new_val != cell.value:
                    cell.value = new_val
                    replaced += 1

    if replaced:
        logger.info("Excel : %s cellule(s) mise(s) à jour", replaced)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
