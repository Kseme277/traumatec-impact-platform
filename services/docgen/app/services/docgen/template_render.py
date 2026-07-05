"""Rendu docxtpl + remplacements texte pour les programmes séminaires."""

from __future__ import annotations

import logging
from datetime import date, datetime
from io import BytesIO
from typing import Any

from docx import Document
from docxtpl import DocxTemplate

from app.services.docgen.highlight_replace import apply_highlight_replacements_docx

logger = logging.getLogger(__name__)


def _format_date_fr(value: str | date | None) -> str:
    if value is None:
        return ""
    if isinstance(value, date):
        d = value
    else:
        text = str(value).strip()[:10]
        try:
            d = date.fromisoformat(text)
        except ValueError:
            return str(value)
    return d.strftime("%d/%m/%Y")


def _format_date_long_fr(value: str | date | None) -> str:
    if value is None:
        return ""
    if isinstance(value, date):
        d = value
    else:
        text = str(value).strip()[:10]
        try:
            d = date.fromisoformat(text)
        except ValueError:
            return str(value)
    months = (
        "janvier",
        "février",
        "mars",
        "avril",
        "mai",
        "juin",
        "juillet",
        "août",
        "septembre",
        "octobre",
        "novembre",
        "décembre",
    )
    return f"{d.day} {months[d.month - 1]} {d.year}"


def _contact_line(event: dict[str, Any]) -> str:
    email = (event.get("national_responsible_email") or event.get("responsible_email") or "").strip()
    phone = (event.get("responsible_phone") or "").strip()
    if email and phone:
        return f"Courriel: {email}        Téléphone: {phone}"
    if email:
        return f"Courriel: {email}"
    if phone:
        return f"Téléphone: {phone}"
    return ""


def build_event_context(event: dict[str, Any]) -> dict[str, Any]:
    """Contexte Jinja2 / docxtpl — alias FR et EN pour les modèles AO."""
    start = event.get("start_date")
    end = event.get("end_date") or start
    from tip_common.location_fields import resolve_lieu_display, resolve_lieu_doc_display

    city = (event.get("city") or "").strip()
    country = (event.get("country") or "").strip()
    lieu = resolve_lieu_display(event)
    lieu_doc = resolve_lieu_doc_display(event)

    lieu_complet = (event.get("lieu_complet") or "").strip() or lieu
    lieu_display = (event.get("lieu_formatted") or lieu_doc or lieu_complet).strip()
    title_display = (event.get("title_formatted") or event.get("title") or "").strip()
    resp_display = (
        event.get("responsible_formatted") or event.get("responsible_person") or ""
    ).strip()
    ctx: dict[str, Any] = {
        "project_number": event.get("project_number") or "",
        "numero_projet": event.get("project_number") or "",
        "title": title_display,
        "titre": title_display,
        "title_formatted": title_display,
        "event_title": title_display,
        "event_name": title_display,
        "event_type": event.get("event_type") or "",
        "type_evenement": event.get("event_type") or "",
        "package_type": event.get("package_type") or "",
        "package_label": event.get("package_label") or "",
        "package_duration_days": event.get("package_duration_days") or event.get("expected_package_days") or 1,
        "activity_label": event.get("activity_label") or "",
        "raw_title": (event.get("raw_title") or event.get("title") or "").strip(),
        "metadata_json": event.get("metadata_json"),
        "city": city,
        "ville": city,
        "lieu": lieu_display,
        "lieu_formatted": lieu_display,
        "location": lieu_display,
        "header_lieu_date": (event.get("header_lieu_date") or "").strip(),
        "date_range_formatted": (event.get("date_range_formatted") or "").strip(),
        "date_single_formatted": (event.get("date_single_formatted") or "").strip(),
        "country": country,
        "pays": country,
        "region": event.get("region") or "",
        "responsible_person": resp_display,
        "responsable": resp_display,
        "responsible_formatted": resp_display,
        "responsible_email": (
            event.get("national_responsible_email") or event.get("responsible_email") or ""
        ).strip(),
        "responsible_phone": (event.get("responsible_phone") or "").strip(),
        "participants_expected": event.get("participants_expected"),
        "participants_count": (
            str(int(event["participants_expected"]))
            if event.get("participants_expected") is not None
            else ""
        ),
        "prepared_by": (event.get("prepared_by") or "").strip(),
        "prepared_by_name": (event.get("prepared_by_name") or event.get("prepared_by") or "").strip(),
        "responsible_photo_bytes": event.get("responsible_photo_bytes"),
        "contact_line": _contact_line(event),
        "preparation_theme": event.get("preparation_theme") or "",
        "theme": event.get("preparation_theme") or "",
        "start_date": _format_date_fr(start),
        "date_debut": _format_date_fr(start),
        "end_date": _format_date_fr(end),
        "date_fin": _format_date_fr(end),
        "event_date": _format_date_fr(start),
        "date_evenement": _format_date_fr(start),
        "start_date_long": _format_date_long_fr(start),
        "end_date_long": _format_date_long_fr(end),
        "year": str(start)[:4] if start else str(datetime.now().year),
        "annee": str(start)[:4] if start else str(datetime.now().year),
        "start_date_raw": str(start)[:10] if start else "",
        "end_date_raw": str(end)[:10] if end else "",
        "cost_center": (event.get("cost_center") or "").strip(),
        "date_du_jour": (event.get("date_single_formatted") or _format_date_long_fr(start) or "").strip(),
        "today": (event.get("date_single_formatted") or _format_date_long_fr(start) or "").strip(),
        "organizer_responsible_name": (event.get("organizer_responsible_name") or "").strip(),
    }
    teacher_names: list[str] = list(event.get("teacher_names") or [])
    if not teacher_names:
        for teacher in event.get("teachers") or []:
            if not isinstance(teacher, dict):
                continue
            name = f"{teacher.get('first_name', '')} {teacher.get('last_name', '')}".strip()
            if name:
                teacher_names.append(name)
    ctx["teacher_names"] = teacher_names
    ctx["teachers"] = event.get("teachers") or []
    for index, name in enumerate(teacher_names[:6], start=1):
        ctx[f"teacher_{index}"] = name
        ctx[f"enseignant_{index}"] = name
    return ctx


def _plain_replacements(context: dict[str, Any]) -> list[tuple[str, str]]:
    """Remplacements texte brut pour modèles sans variables Jinja."""
    from tip_common.french_placeholders import build_french_placeholder_pairs

    pairs: list[tuple[str, str]] = []
    mapping = {
        "Zurich": context.get("city") or "Zurich",
        "{{ project_number }}": context.get("project_number", ""),
        "{{ title }}": context.get("title", ""),
        "{{ city }}": context.get("city", ""),
        "{{ country }}": context.get("country", ""),
        "{{ start_date }}": context.get("start_date", ""),
        "{{ end_date }}": context.get("end_date", ""),
        "[PROJECT_NUMBER]": context.get("project_number", ""),
        "[EVENT_TITLE]": context.get("title", ""),
        "[CITY]": context.get("city", ""),
        "[COUNTRY]": context.get("country", ""),
        "[START_DATE]": context.get("start_date", ""),
        "[END_DATE]": context.get("end_date", ""),
        "[RESPONSIBLE]": context.get("responsible_person", ""),
    }
    for old, new in mapping.items():
        if new and old != new:
            pairs.append((old, str(new)))
    seen = {old for old, _ in pairs}
    for old, new in build_french_placeholder_pairs(context):
        if old not in seen:
            pairs.append((old, new))
            seen.add(old)
    return pairs


def _replace_in_paragraph(paragraph, replacements: list[tuple[str, str]]) -> None:
    for old, new in replacements:
        if old in paragraph.text:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)


def _apply_plain_replacements_docx(data: bytes, replacements: list[tuple[str, str]]) -> bytes:
    if not replacements:
        return data
    doc = Document(BytesIO(data))
    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, replacements)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def render_package_document(
    template_bytes: bytes,
    *,
    suffix: str,
    context: dict[str, Any],
    replacement_fields: list[dict[str, Any]] | None = None,
    document_role: str | None = None,
    day_index: int | None = None,
) -> bytes:
    """Fusionne le template avec les données événement et remplace TOUTES les sections détectées."""
    from app.services.docgen.comprehensive_replace import (
        apply_exhaustive_pairs_doc_binary,
        apply_exhaustive_pairs_docx,
        apply_exhaustive_pairs_xlsx,
        build_exhaustive_pairs,
    )
    from app.services.docgen.document_role_replace import _date_for_role
    from app.services.docgen.docx_xml_replace import apply_role_replacements_docx
    from app.services.docgen.excel_replace import apply_excel_replacements
    from app.services.docgen.legacy_doc_replace import apply_legacy_doc_replacements
    from tip_common.location_fields import resolve_lieu_doc_display

    context = dict(context)
    if day_index and document_role in {"presence_enseignants", "presence_participants"}:
        context["date_du_jour"] = _date_for_role(context, document_role, day_index)
    lieu_doc = resolve_lieu_doc_display(context)
    if lieu_doc:
        context["lieu_formatted"] = lieu_doc
        context["lieu"] = lieu_doc
        context["location"] = lieu_doc

    fields = replacement_fields or []
    exhaustive = build_exhaustive_pairs(context, fields)
    programme_like_roles = {
        "programme",
        "presence_enseignants",
        "presence_participants",
        "liste_definitive",
        "badge",
    }
    is_programme_like = document_role in programme_like_roles
    strict_roles = {
        "budget",
        "coordonnees_bancaires",
        "rapport_national",
        "accuse_paiement",
        "accord_collaboration",
        "rapport_depenses",
    }

    ext = suffix.lower()
    if document_role in {"coordonnees_bancaires", "accuse_paiement"} and ext == ".docx":
        from app.services.docgen.coordonnees_docx_replace import apply_coordonnees_docx_replacements

        return apply_coordonnees_docx_replacements(
            template_bytes,
            context,
            replacement_fields=fields,
            document_role=document_role,
        )

    if document_role == "rapport_national" and ext == ".docx":
        from app.services.docgen.rapport_national_docx_replace import (
            apply_rapport_national_docx_replacements,
        )

        return apply_rapport_national_docx_replacements(
            template_bytes,
            context,
            replacement_fields=fields,
        )

    if ext == ".xlsx":
        rendered = apply_excel_replacements(
            template_bytes,
            context,
            replacement_fields=fields,
            document_role=document_role,
        )
        # Budget : remplacements dédiés uniquement. Autres xlsx : placeholders {{…}} restants.
        if document_role == "budget":
            return rendered
        return apply_exhaustive_pairs_xlsx(rendered, exhaustive)
    if ext == ".doc":
        if is_programme_like:
            from app.services.docgen.programme_doc_replace import apply_programme_doc_replacements

            return apply_programme_doc_replacements(
                template_bytes,
                context,
                replacement_fields=fields,
            )
        rendered = apply_legacy_doc_replacements(
            template_bytes,
            context,
            replacement_fields=fields,
            document_role=document_role,
        )
        if document_role in strict_roles or is_programme_like:
            return rendered
        return apply_exhaustive_pairs_doc_binary(rendered, exhaustive)
    if ext != ".docx":
        logger.warning("Format %s : copie sans fusion automatique", ext)
        return template_bytes

    if is_programme_like:
        # Les programmes AO utilisent les zones surlignées, pas des variables Jinja partout.
        rendered = template_bytes
    else:
        try:
            doc = DocxTemplate(BytesIO(template_bytes))
            doc.render(context)
            buffer = BytesIO()
            doc.save(buffer)
            rendered = buffer.getvalue()
        except Exception as exc:
            logger.warning("docxtpl render fallback (%s), remplacements texte seuls", exc)
            rendered = template_bytes

    replacements = (
        _plain_replacements(context)
        if document_role not in strict_roles and not is_programme_like
        else []
    )
    if document_role not in strict_roles and not is_programme_like:
        for old, new in exhaustive:
            if (old, new) not in replacements:
                replacements.append((old, new))
    if replacements:
        rendered = _apply_plain_replacements_docx(rendered, replacements)
    rendered = apply_role_replacements_docx(
        rendered,
        context,
        document_role=document_role,
        day_index=day_index,
        replacement_fields=fields,
    )
    highlight_fields = [] if document_role in strict_roles else fields
    rendered = apply_highlight_replacements_docx(
        rendered, context, replacement_fields=highlight_fields
    )
    if is_programme_like and document_role not in {
        "presence_enseignants",
        "presence_participants",
    }:
        rendered = apply_exhaustive_pairs_docx(rendered, exhaustive)
    if document_role in strict_roles or is_programme_like:
        if document_role == "programme" and context.get("responsible_photo_bytes"):
            from app.services.docgen.docx_photo_replace import apply_responsible_photo_docx

            rendered = apply_responsible_photo_docx(
                rendered,
                context.get("responsible_photo_bytes"),
                document_role=document_role,
            )
        return rendered
    return apply_exhaustive_pairs_docx(rendered, exhaustive)


def render_programme_document(
    template_bytes: bytes,
    *,
    suffix: str,
    context: dict[str, Any],
) -> bytes:
    return render_package_document(template_bytes, suffix=suffix, context=context)
