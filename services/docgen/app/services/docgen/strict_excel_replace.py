"""Remplacements Excel ciblés pour formulaires stricts (rapport dépenses, …)."""

from __future__ import annotations

import logging
import re
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.docgen.document_role_replace import _strict_fields_for_role
from app.services.docgen.excel_cell_utils import set_cell_value

logger = logging.getLogger(__name__)

_PREPARED_BY_COLON_RE = re.compile(r'^(Prepared by:\s*"?)(.+?)("?)$', re.I)
_PREPARED_BY_PLAIN_RE = re.compile(r"^(Prepared by\s+)(.+)$", re.I)


def _replace_prepared_by_cell(value: str, context: dict[str, Any]) -> str | None:
    prepared = (context.get("prepared_by_name") or context.get("prepared_by") or "").strip()
    if not prepared:
        return None
    stripped = value.strip()
    m = _PREPARED_BY_COLON_RE.match(stripped)
    if m and m.group(2).strip() != prepared:
        return f"{m.group(1)}{prepared}{m.group(3)}"
    m2 = _PREPARED_BY_PLAIN_RE.match(stripped)
    if m2 and m2.group(2).strip() != prepared:
        return f"{m2.group(1)}{prepared}"
    return None


def apply_strict_excel_replacements(
    xlsx_bytes: bytes,
    context: dict[str, Any],
    *,
    replacement_fields: list[dict[str, Any]] | None = None,
    document_role: str | None = None,
) -> bytes:
    """Remplacements exacts (échantillons analysés) + champs métier ciblés."""
    from tip_common.template_field_analyzer import context_value_for_key, format_value_for_field

    fields = _strict_fields_for_role(replacement_fields)
    extra_prepared_by = document_role == "rapport_depenses"

    if not fields and not extra_prepared_by:
        return xlsx_bytes

    try:
        wb = load_workbook(BytesIO(xlsx_bytes))
    except Exception as exc:
        logger.warning("Excel strict illisible (%s)", exc)
        return xlsx_bytes

    replaced = 0
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or not cell.value.strip():
                    continue
                stripped = cell.value.strip()
                updated: str | None = None

                for field in fields:
                    sample = str(field.get("sample", "")).strip()
                    if not sample or stripped != sample:
                        continue
                    value = format_value_for_field(field, context) or context_value_for_key(
                        str(field.get("context_key", "")), context
                    )
                    if value and value != stripped:
                        updated = value
                    break

                if updated is None and extra_prepared_by:
                    updated = _replace_prepared_by_cell(stripped, context)

                if updated is None and "{{" in stripped:
                    from tip_common.french_placeholders import build_french_placeholder_pairs

                    candidate = stripped
                    for old, new in sorted(
                        build_french_placeholder_pairs(context),
                        key=lambda item: -len(item[0]),
                    ):
                        if old in candidate:
                            candidate = candidate.replace(old, new)
                    if candidate != stripped:
                        updated = candidate

                if updated is not None and updated != stripped:
                    set_cell_value(sheet, cell, updated)
                    replaced += 1

    if not replaced:
        return xlsx_bytes

    logger.info(
        "Excel strict (%s) : %s cellule(s) mise(s) à jour",
        document_role or "?",
        replaced,
    )
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
