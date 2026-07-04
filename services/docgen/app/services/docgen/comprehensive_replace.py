"""Remplacement exhaustif : toutes les paires sample→valeur dans tout le document."""

from __future__ import annotations

import logging
import re
import zipfile
from io import BytesIO
from typing import Any
from xml.etree import ElementTree as ET

from app.services.docgen.excel_cell_utils import set_cell_value

logger = logging.getLogger(__name__)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _paragraph_text(paragraph: ET.Element) -> str:
    return "".join(
        (node.text or "")
        for run in paragraph.iter(f"{{{W_NS}}}r")
        for node in run.iter(f"{{{W_NS}}}t")
    )


def _set_paragraph_text(paragraph: ET.Element, new_text: str) -> None:
    nodes = [
        node
        for run in paragraph.iter(f"{{{W_NS}}}r")
        for node in run.iter(f"{{{W_NS}}}t")
    ]
    if not nodes:
        return
    nodes[0].text = new_text
    for node in nodes[1:]:
        node.text = ""


def _context_fallback_pairs(context: dict[str, Any]) -> list[tuple[str, str]]:
    """Paires génériques présentes dans la plupart des modèles AO."""
    from tip_common.french_placeholders import build_french_placeholder_pairs

    pairs: list[tuple[str, str]] = []
    mapping = {
        "TBD": context.get("city") or context.get("lieu") or "",
        "Zurich": context.get("city") or "",
        "[PROJECT_NUMBER]": context.get("project_number") or "",
        "[EVENT_TITLE]": context.get("title") or "",
        "[CITY]": context.get("city") or "",
        "[COUNTRY]": context.get("country") or "",
        "[START_DATE]": context.get("start_date") or "",
        "[END_DATE]": context.get("end_date") or "",
        "[RESPONSIBLE]": context.get("responsible_person") or "",
        "adresse@email": context.get("responsible_email") or "",
        "{{ project_number }}": context.get("project_number") or "",
        "{{ title }}": context.get("title") or "",
        "{{ city }}": context.get("city") or "",
        "{{ country }}": context.get("country") or "",
        "{{ start_date }}": context.get("start_date") or "",
        "{{ end_date }}": context.get("end_date") or "",
    }
    for old, new in mapping.items():
        if new and old != new:
            pairs.append((old, str(new)))
    seen = {old for old, _ in pairs}
    for old, new in build_french_placeholder_pairs(context):
        if old not in seen:
            pairs.append((old, new))
            seen.add(old)
    return pairs


def build_exhaustive_pairs(
    context: dict[str, Any],
    replacement_fields: list[dict[str, Any]] | None,
) -> list[tuple[str, str]]:
    """Toutes les paires de remplacement triées (plus long d'abord)."""
    from tip_common.template_field_analyzer import build_replacement_pairs

    pairs = build_replacement_pairs(replacement_fields, context)
    pairs.extend(_context_fallback_pairs(context))
    seen: set[str] = set()
    unique: list[tuple[str, str]] = []
    for old, new in sorted(pairs, key=lambda x: -len(x[0])):
        if old in seen or old == new or not new:
            continue
        seen.add(old)
        unique.append((old, new))
    return unique


def apply_exhaustive_pairs_docx(data: bytes, pairs: list[tuple[str, str]]) -> bytes:
    """Remplace chaque sample dans tous les paragraphes Word (corps + en-têtes)."""
    if not pairs:
        return data
    try:
        with zipfile.ZipFile(BytesIO(data), "r") as zin:
            parts = {name: zin.read(name) for name in zin.namelist()}
    except zipfile.BadZipFile:
        return data

    total = 0
    for part_name, content in list(parts.items()):
        if not part_name.startswith("word/") or not part_name.endswith(".xml"):
            continue
        try:
            root = ET.fromstring(content)
        except ET.ParseError:
            continue
        part_count = 0
        for paragraph in root.iter(f"{{{W_NS}}}p"):
            original = _paragraph_text(paragraph)
            if not original.strip():
                continue
            updated = original
            for old, new in pairs:
                if old in updated:
                    updated = updated.replace(old, new)
            if updated != original:
                _set_paragraph_text(paragraph, updated)
                part_count += 1
        if part_count:
            parts[part_name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            total += part_count

    if not total:
        return data

    from app.services.docgen.docx_zip_repack import repack_docx_archive

    logger.info("Remplacement exhaustif docx : %s paragraphe(s)", total)
    return repack_docx_archive(data, parts)


def apply_exhaustive_pairs_doc_binary(data: bytes, pairs: list[tuple[str, str]]) -> bytes:
    from app.services.docgen.doc_binary_replace import replace_fixed_width_in_binary

    if not pairs:
        return data
    result = data
    replaced = 0
    for old, new in pairs:
        result, count = replace_fixed_width_in_binary(result, old, new)
        replaced += count
    if replaced:
        logger.info("Remplacement exhaustif .doc : %s chaîne(s)", replaced)
    return result


def apply_exhaustive_pairs_xlsx(data: bytes, pairs: list[tuple[str, str]]) -> bytes:
    if not pairs:
        return data
    try:
        from openpyxl import load_workbook
    except ImportError:
        return data

    try:
        wb = load_workbook(BytesIO(data))
    except Exception:
        return data

    replaced = 0
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or not cell.value.strip():
                    continue
                updated = cell.value
                for old, new in pairs:
                    if old in updated:
                        updated = updated.replace(old, new)
                if updated != cell.value and set_cell_value(sheet, cell, updated):
                    replaced += 1

    if not replaced:
        return data

    out = BytesIO()
    wb.save(out)
    logger.info("Remplacement exhaustif xlsx : %s cellule(s)", replaced)
    return out.getvalue()
