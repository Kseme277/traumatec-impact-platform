from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from collections.abc import Callable

from openpyxl import load_workbook

REQUIRED_FIELDS = ("project_number", "title")

# Format réel AO Alliance — Projects.xlsx (export Traumatec / TIP)
PROJECTS_XLSX_COLUMNS = (
    "Title",
    "Activity",
    "Project number",
    "Start date",
    "End date",
    "Status",
    "Responsible person",
    "Organizer responsible",
    "Organizer email",
    "Location",
    "Country",
    "Region",
    "Cost center",
    "Participants (expected nb)",
    "Participants (real nb)",
    "Amount (CHF)",
    "Payments done (CHF)",
    "% paid",
    "Balance to pay (CHF)",
)

EXPECTED_COLUMNS_HELP = (
    "Format Projects.xlsx AO Alliance : colonnes « Title », « Project number », "
    "« Activity », « Start date », « End date », « Location », « Country », etc."
)

HEADER_ALIASES: dict[str, str] = {
    # Projects.xlsx
    "title": "title",
    "project title": "title",
    "project number": "project_number",
    "activity": "event_type",
    "start date": "start_date",
    "end date": "end_date",
    "status": "project_status",
    "responsible person": "responsible_person",
    "responsable": "responsible_person",
    "organizer responsible": "organizer_responsible_name",
    "responsable organisation": "organizer_responsible_name",
    "responsable organisateur": "organizer_responsible_name",
    "support administratif": "organizer_responsible_name",
    "organizer email": "organizer_responsible_email",
    "email responsable organisation": "organizer_responsible_email",
    "email support": "organizer_responsible_email",
    "location": "city",
    "country": "country",
    "region": "region",
    "cost center": "cost_center",
    "participants expected nb": "participants_expected",
    "participants real nb": "participants_real",
    "amount chf": "amount_chf",
    "payments done chf": "payments_done_chf",
    "paid": "percent_paid",
    "balance to pay chf": "balance_to_pay_chf",
    # Variantes FR / autres exports
    "project no": "project_number",
    "project no.": "project_number",
    "project nr": "project_number",
    "numero projet": "project_number",
    "numero de projet": "project_number",
    "n projet": "project_number",
    "no projet": "project_number",
    "event name": "title",
    "event title": "title",
    "titre": "title",
    "nom evenement": "title",
    "event type": "event_type",
    "type evenement": "event_type",
    "preparation theme": "preparation_theme",
    "theme": "preparation_theme",
    "pays": "country",
    "ville": "city",
    "city": "city",
    "date debut": "start_date",
    "date fin": "end_date",
}

FIELD_PATTERNS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("project_number", ("project number", "numero projet", "numero de projet")),
    ("title", ("project title", "event name", "event title", "nom evenement")),
    ("event_type", ("event type", "type evenement", "activity")),
    ("preparation_theme", ("preparation theme", "theme preparation")),
    ("country", ("country name",)),
    ("city", ("location city",)),
    ("start_date", ("start date", "date debut")),
    ("end_date", ("end date", "date fin")),
    ("project_status", ("status", "statut projet")),
    ("responsible_person", ("responsible person", "responsable", "person in charge")),
    ("organizer_responsible_name", ("organizer responsible", "responsable organisation", "support administratif")),
    ("organizer_responsible_email", ("organizer email", "email responsable organisation", "email support")),
    ("region", ("region",)),
    ("cost_center", ("cost center",)),
    ("participants_expected", ("participants expected", "expected nb")),
    ("participants_real", ("participants real", "real nb")),
    ("amount_chf", ("amount chf",)),
    ("payments_done_chf", ("payments done",)),
    ("percent_paid", ("% paid", "percent paid")),
    ("balance_to_pay_chf", ("balance to pay",)),
)

NUMERIC_FIELDS = frozenset(
    {"amount_chf", "payments_done_chf", "percent_paid", "balance_to_pay_chf"},
)
INTEGER_FIELDS = frozenset({"participants_expected", "participants_real"})

THEME_ALIASES = {
    "operatoire": "operatory",
    "operatory": "operatory",
    "oper": "operatory",
    "op c": "operatory",
    "nonop": "operatory",
    "non-op": "operatory",
    "cmf": "operatory",
    "orp": "pbo",
    "orp s": "pbo",
    "orp c": "pbo",
    "pbo": "pbo",
    "iec": "iec",
    "iec s": "iec",
}

FIELD_MAX_LENGTHS: dict[str, int] = {
    "project_number": 64,
    "title": 512,
    "event_type": 255,
    "country": 128,
    "city": 128,
    "region": 128,
    "responsible_person": 255,
    "organizer_responsible_name": 255,
    "organizer_responsible_email": 255,
    "project_status": 128,
    "cost_center": 255,
}


def _strip_accents(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text)
    return "".join(char for char in normalized if not unicodedata.combining(char))


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = _strip_accents(text)
    text = text.replace("_", " ")
    text = text.replace("/", " ")
    text = text.replace("-", " ")
    text = re.sub(r"[^\w\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _resolve_field(normalized: str) -> str | None:
    if not normalized:
        return None

    direct = HEADER_ALIASES.get(normalized)
    if direct:
        return direct

    best_field: str | None = None
    best_score = 0
    for field, patterns in FIELD_PATTERNS:
        for pattern in patterns:
            if normalized == pattern:
                return field
            if pattern in normalized:
                score = len(pattern)
                if score > best_score:
                    best_score = score
                    best_field = field
    return best_field


def _build_column_map(header_row: tuple[Any, ...]) -> dict[int, str]:
    column_map: dict[int, str] = {}
    for index, header in enumerate(header_row):
        field = _resolve_field(_normalize_header(header))
        if field and field not in column_map.values():
            column_map[index] = field
    return column_map


def _find_header_row(rows: list[tuple[Any, ...]]) -> tuple[int, dict[int, str]] | None:
    for row_index, row in enumerate(rows[:25]):
        column_map = _build_column_map(row)
        if all(field in column_map.values() for field in REQUIRED_FIELDS):
            return row_index, column_map
    return None


def _format_detected_headers(rows: list[tuple[Any, ...]]) -> str:
    snippets: list[str] = []
    for row_index, row in enumerate(rows[:5], start=1):
        labels = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
        if labels:
            snippets.append(f"L{row_index}: {', '.join(labels[:12])}")
    return " | ".join(snippets) if snippets else "aucune colonne lisible"


def _parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # Numéro de série Excel (openpyxl peut renvoyer un float)
        try:
            from openpyxl.utils.datetime import from_excel

            parsed = from_excel(value)
            if isinstance(parsed, datetime):
                return parsed.date()
            if isinstance(parsed, date):
                return parsed
        except (ValueError, TypeError, OverflowError):
            pass
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _infer_theme_from_text(value: Any) -> str | None:
    if value is None or value == "":
        return None
    text = _strip_accents(str(value).strip().lower())
    for alias, canonical in THEME_ALIASES.items():
        if alias in text:
            return canonical
    return None


def _header_label(value: Any, index: int) -> str:
    text = str(value).strip() if value is not None else ""
    return text or f"Column_{index + 1}"


def _build_header_labels(header_row: tuple[Any, ...]) -> list[str]:
    labels: list[str] = []
    seen: dict[str, int] = {}
    for index, cell in enumerate(header_row):
        label = _header_label(cell, index)
        count = seen.get(label, 0)
        seen[label] = count + 1
        if count:
            label = f"{label} ({count + 1})"
        labels.append(label)
    return labels


def _serialize_excel_value(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, Decimal):
        return float(value)
    return value


def _row_excel_snapshot(header_labels: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    snapshot: dict[str, Any] = {}
    for index, label in enumerate(header_labels):
        if index >= len(row):
            snapshot[label] = None
            continue
        snapshot[label] = _serialize_excel_value(row[index])
    return snapshot


def _parse_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    text = str(value).strip().replace(" ", "")
    if not text:
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        return None


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(" ", "").replace("'", "")
    if not text:
        return None
    if text.endswith("%"):
        text = text[:-1]
    text = text.replace(",", "")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _parse_percent(value: Any) -> Decimal | None:
    parsed = _parse_decimal(value)
    if parsed is None:
        return None
    if parsed > 1:
        return (parsed / Decimal(100)).quantize(Decimal("0.0001"))
    return parsed


def _assign_typed_field(data: dict[str, Any], field: str, value: Any) -> None:
    if field in INTEGER_FIELDS:
        data[field] = _parse_int(value)
    elif field == "percent_paid":
        data[field] = _parse_percent(value)
    elif field in NUMERIC_FIELDS:
        data[field] = _parse_decimal(value)
    elif field in {"start_date", "end_date"}:
        data[field] = _parse_date(value)
    elif field == "preparation_theme":
        data[field] = _infer_theme_from_text(value)
    else:
        data[field] = _cell_str(value)


def _cell_str(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _fit_text(
    value: str | None,
    field: str,
    row_index: int,
    errors: list[str],
) -> str | None:
    if value is None:
        return None
    limit = FIELD_MAX_LENGTHS.get(field)
    if limit and len(value) > limit:
        errors.append(f"Ligne {row_index} : champ « {field} » tronqué ({len(value)} → {limit} car.)")
        return value[:limit]
    return value


def _normalize_event_row(data: dict[str, Any], row_index: int, errors: list[str]) -> dict[str, Any]:
    for field in FIELD_MAX_LENGTHS:
        if field in data and isinstance(data[field], str):
            data[field] = _fit_text(data[field], field, row_index, errors)
    return data


def _load_rows(content: bytes) -> list[tuple[Any, ...]]:
    """Charge toutes les colonnes — read_only ne lit qu'une colonne sur Projects.xlsx."""
    workbook = load_workbook(filename=BytesIO(content), data_only=True)
    sheet = workbook.active
    max_row = sheet.max_row
    max_col = sheet.max_column
    rows: list[tuple[Any, ...]] = []
    for row_index in range(1, max_row + 1):
        rows.append(tuple(sheet.cell(row_index, col).value for col in range(1, max_col + 1)))
    workbook.close()
    return rows


def parse_projects_workbook(
    content: bytes,
    on_progress: Callable[[int, int], None] | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    """Parse Projects.xlsx AO Alliance (Title, Project number, Activity, dates, Location, Country…)."""
    rows = _load_rows(content)
    if not rows:
        return [], ["Fichier Excel vide"]

    header_match = _find_header_row(rows)
    if not header_match:
        return [], [
            "En-têtes requis introuvables. "
            f"{EXPECTED_COLUMNS_HELP} "
            f"Détecté : {_format_detected_headers(rows)}"
        ]

    header_row_index, column_map = header_match
    header_labels = _build_header_labels(rows[header_row_index])
    data_rows = rows[header_row_index + 1 :]
    total_data_rows = len(data_rows)
    events: list[dict[str, Any]] = []
    errors: list[str] = []

    for offset, row in enumerate(data_rows):
        row_index = header_row_index + 2 + offset
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        excel_snapshot = _row_excel_snapshot(header_labels, row)
        data: dict[str, Any] = {"metadata_json": {"excel": excel_snapshot, "source": "projects.xlsx"}}
        for col_index, field in column_map.items():
            if col_index >= len(row):
                continue
            _assign_typed_field(data, field, row[col_index])

        if not data.get("preparation_theme"):
            from tip_common.package_types import infer_preparation_theme_for_event

            data["preparation_theme"] = infer_preparation_theme_for_event(
                event_type=data.get("event_type"),
                title=data.get("title"),
            ) or _infer_theme_from_text(data.get("event_type"))

        project_number = data.get("project_number")
        title = data.get("title")
        if not project_number or not title:
            errors.append(f"Ligne {row_index} ignorée : numéro projet ou titre manquant")
            continue

        events.append(_normalize_event_row(data, row_index, errors))

        if on_progress and total_data_rows and (offset + 1) % 100 == 0:
            on_progress(offset + 1, total_data_rows)

    if on_progress and total_data_rows:
        on_progress(total_data_rows, total_data_rows)

    return events, errors
