from __future__ import annotations

import re
import unicodedata
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

REGISTRATION_COLUMNS_HELP = (
    "Export plateforme d'inscription participants : "
    "Nom, prenom, Statut, Nom_evenement, Formation_sanitaire, Email, etc."
)

_HEADER_ALIASES: dict[str, str] = {
    "nom": "last_name",
    "prenom": "first_name",
    "prénom": "first_name",
    "statut": "statut",
    "nom_evenement": "event_title_source",
    "nom_événement": "event_title_source",
    "pays_evenement": "country",
    "formation_sanitaire": "hospital",
    "email": "email",
    "telephone": "phone",
    "téléphone": "phone",
    "tel": "phone",
    "mobile": "phone",
    "gsm": "phone",
    "portable": "phone",
    "numero_telephone": "phone",
    "numero_de_telephone": "phone",
    "numéro_téléphone": "phone",
    "numero_tel": "phone",
    "submissionid": "submission_id",
    "specialite": "specialite",
    "autreprofil": "autreprofil",
}

_ENSEIGNANT_KEYWORDS = (
    "enseignant",
    "formateur",
    "instructor",
    "faculty",
    "teacher",
    "trainer",
)

_FRENCH_MONTHS: dict[str, int] = {
    "janvier": 1,
    "janv": 1,
    "fevrier": 2,
    "février": 2,
    "fev": 2,
    "fév": 2,
    "mars": 3,
    "avril": 4,
    "avr": 4,
    "mai": 5,
    "juin": 6,
    "juillet": 7,
    "juil": 7,
    "aout": 8,
    "août": 8,
    "septembre": 9,
    "sept": 9,
    "octobre": 10,
    "oct": 10,
    "novembre": 11,
    "nov": 11,
    "decembre": 12,
    "décembre": 12,
    "dec": 12,
    "déc": 12,
}

_COUNTRY_ALIASES: dict[str, set[str]] = {
    "cameroun": {"cameroun", "cameroon"},
    "cameroon": {"cameroun", "cameroon"},
}


def _normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "_", text.strip().lower()).strip("_")


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_token(value: str) -> str:
    text = unicodedata.normalize("NFKD", value.lower())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", text)


def _detect_certificate_role(statut: str, specialite: str, autreprofil: str) -> str:
    blob = " ".join((statut, specialite, autreprofil)).lower()
    if any(key in blob for key in _ENSEIGNANT_KEYWORDS):
        return "enseignant"
    return "participant"


def _build_full_name(last_name: str, first_name: str) -> str:
    parts = [p for p in (last_name, first_name) if p]
    return " ".join(parts) if parts else "Participant"


def _titles_match(a: str, b: str) -> bool:
    na, nb = _normalize_token(a), _normalize_token(b)
    if not na or not nb:
        return True
    return na in nb or nb in na or _similar_ratio(na, nb) >= 0.55


def _similar_ratio(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    shorter, longer = (a, b) if len(a) <= len(b) else (b, a)
    matches = sum(1 for i, ch in enumerate(shorter) if i < len(longer) and longer[i] == ch)
    return matches / max(len(longer), 1)


def _parse_iso_date(value: str | date | None) -> date | None:
    if value is None:
        return None
    if isinstance(value, date):
        return value
    text = str(value).strip()[:10]
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _extract_registration_date_hint(title: str) -> tuple[int, int] | None:
    """Extrait jour/mois du préfixe inscription (« 13 juin, Séminaire… »)."""
    text = unicodedata.normalize("NFKD", title.strip().lower())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    match = re.match(r"^(\d{1,2})\s+([a-z]+)", text)
    if not match:
        return None
    day = int(match.group(1))
    month_name = match.group(2).rstrip(",")
    month = _FRENCH_MONTHS.get(month_name)
    if month is None or not 1 <= day <= 31:
        return None
    return day, month


def _country_tokens(country: str | None) -> set[str]:
    if not country:
        return set()
    token = _normalize_token(country)
    if not token:
        return set()
    aliases = set()
    for key, group in _COUNTRY_ALIASES.items():
        if token == _normalize_token(key) or token in {_normalize_token(v) for v in group}:
            aliases |= {_normalize_token(v) for v in group}
    return {token, *aliases}


def _location_signals_match(
    source_title: str,
    pays_evenement: str,
    *,
    city: str | None,
    country: str | None,
) -> bool:
    """Vérifie que ville/pays de l'export concordent avec l'événement TIP."""
    blob = _normalize_token(f"{source_title} {pays_evenement}")
    if not blob:
        return True

    city_token = _normalize_token(city or "")
    country_tokens = _country_tokens(country)
    pays_tokens = _country_tokens(pays_evenement)

    city_ok = not city_token or city_token in blob
    country_ok = True
    if country_tokens or pays_tokens:
        all_country = country_tokens | pays_tokens
        country_ok = any(token in blob for token in all_country if token)

    return city_ok and country_ok


def _registration_row_matches_event(
    record: dict[str, Any],
    expected_event: dict[str, Any] | None,
) -> bool:
    if not expected_event:
        return True

    source_title = record.get("event_title_source", "")
    pays = record.get("country", "")
    start = _parse_iso_date(expected_event.get("start_date"))
    end = _parse_iso_date(expected_event.get("end_date")) or start

    if start and source_title:
        hint = _extract_registration_date_hint(source_title)
        if hint:
            day, month = hint
            start_ok = start.day == day and start.month == month
            end_ok = end is None or (end.day == day and end.month == month)
            if not (start_ok or end_ok):
                return False

    if source_title or pays:
        if not _location_signals_match(
            source_title,
            pays,
            city=expected_event.get("city"),
            country=expected_event.get("country"),
        ):
            return False

    return True


def _events_are_same_context(
    source_title: str,
    pays_evenement: str,
    expected_event: dict[str, Any] | None,
) -> bool:
    """Même événement si date + lieu concordent (titres différents acceptés)."""
    if not expected_event:
        return True
    fake_record = {"event_title_source": source_title, "country": pays_evenement}
    return _registration_row_matches_event(fake_record, expected_event)


def parse_registration_workbook(
    content: bytes,
    *,
    expected_event: dict[str, Any] | None = None,
    expected_event_title: str | None = None,
) -> tuple[list[dict[str, Any]], list[str], str | None]:
    """Parse l'export Excel de la plateforme d'inscription participants."""
    if expected_event is None and expected_event_title:
        expected_event = {"title": expected_event_title}

    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        raise ValueError("Fichier Excel vide.")

    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header_row):
        key = _normalize_header(cell)
        field = _HEADER_ALIASES.get(key)
        if field:
            col_map[idx] = field

    if "last_name" not in col_map.values() and "first_name" not in col_map.values():
        raise ValueError(
            f"Colonnes Nom/prenom introuvables. {REGISTRATION_COLUMNS_HELP}"
        )

    parsed: list[dict[str, Any]] = []
    warnings: list[str] = []
    source_titles: set[str] = set()
    skipped_other_event = 0
    all_source_titles: set[str] = set()

    for row_idx, row in enumerate(rows_iter, start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        record: dict[str, Any] = {"row_number": row_idx}
        for col_idx, field in col_map.items():
            if col_idx < len(row):
                record[field] = _normalize_text(row[col_idx])

        last_name = record.get("last_name", "")
        first_name = record.get("first_name", "")
        if not last_name and not first_name:
            continue

        event_title_source = record.get("event_title_source", "")
        if event_title_source:
            all_source_titles.add(event_title_source)

        if not _registration_row_matches_event(record, expected_event):
            skipped_other_event += 1
            continue

        statut = record.get("statut", "")
        specialite = record.get("specialite", "")
        autreprofil = record.get("autreprofil", "")
        role = _detect_certificate_role(statut, specialite, autreprofil)
        if event_title_source:
            source_titles.add(event_title_source)

        parsed.append(
            {
                "full_name": _build_full_name(last_name, first_name),
                "last_name": last_name or None,
                "first_name": first_name or None,
                "hospital": record.get("hospital") or None,
                "email": record.get("email") or None,
                "phone": record.get("phone") or None,
                "statut": statut or None,
                "certificate_role": role,
                "row_number": row_idx,
                "registration_meta": {
                    **{
                        k: v
                        for k, v in record.items()
                        if k
                        not in {
                            "last_name",
                            "first_name",
                            "hospital",
                            "email",
                            "phone",
                            "statut",
                            "row_number",
                        }
                        and v
                    },
                    **({"event_title_source": event_title_source} if event_title_source else {}),
                },
            }
        )

    if skipped_other_event > 0:
        warnings.append(
            f"{skipped_other_event} ligne(s) ignorée(s) — autre date ou lieu dans le fichier d'export."
        )

    unmatched_titles = {
        title
        for title in all_source_titles
        if not _events_are_same_context(title, "", expected_event)
    }
    if len(unmatched_titles) >= 1 and skipped_other_event == 0 and len(all_source_titles) > 1:
        warnings.append(
            "Le fichier contient plusieurs sessions — seules les lignes correspondant "
            "à la date et au lieu de l'événement TIP ont été importées."
        )

    source_title = next(iter(source_titles), None) or next(iter(all_source_titles), None)
    expected_title = (expected_event or {}).get("title") if expected_event else None

    if expected_event and source_title and expected_title:
        context_ok = _events_are_same_context(
            source_title,
            "",
            expected_event,
        )
        title_ok = _titles_match(expected_title, source_title)
        if not context_ok and not title_ok:
            warnings.append(
                "La date ou le lieu de l'export ne correspond pas à l'événement TIP sélectionné. "
                "Vérifiez le fichier ou l'événement choisi."
            )

    if not parsed:
        if skipped_other_event > 0:
            raise ValueError(
                "Aucun participant ne correspond à la date et au lieu de l'événement sélectionné. "
                "Vérifiez que l'export provient bien de la plateforme d'inscription de l'événement."
            )
        raise ValueError("Aucun participant valide trouvé dans le fichier.")

    return parsed, warnings, source_title
